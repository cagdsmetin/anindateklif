"""Zip Perde bayi fiyat tablosu.

Albert Genau'dan farkli olarak burada formul yok: tedarikcinin (SKYART)
Excel'i dogrudan EN x BOY (cm) -> adet fiyati (EUR) izgarasidir. Tablo
merkezi olarak tutulur (Mongo `zip_perde_config`, "id": "default"); yeni
fiyat listesi geldiginde admin ayni formattaki Excel'i tekrar yukler
(bkz. server.py /zip-perde/admin-upload) ve tum Zip Perde bayileri aninda
yeni tabloyu kullanir.

Ara olculer icin fiyat BIR UST basamaktan okunur (EN 210 -> 225 sutunu):
tedarikci de kesimi ust basamaktan fiyatliyor, alt basamak zarar ettirir.
Bu arama istemci tarafinda da (frontend/src/lib/zip-perde.ts) aynen
yapiliyor -- teklif ekraninda her tus vurusunda sunucuya gitmemek icin.
"""
import io
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

CURRENCY = "EUR"

_DATA_PATH = Path(__file__).parent / "data" / "zip_perde_price_data.json"
with open(_DATA_PATH, encoding="utf-8") as _f:
    DEFAULT_TABLE: Dict[str, Any] = json.load(_f)

# Excel'deki hucreler "=ROUND(250*(1+$D$2/100),0)" seklinde: taban fiyat x
# zam orani. Hucrenin onbellekli degeri yoksa (dosya Excel'de hic
# kaydedilmeden uretildiyse) tabani ve zammi formulden kendimiz cozeriz.
_ROUND_RE = re.compile(r"^=\s*ROUND\(\s*([\d.]+)\s*\*")


class ZipPerdeTabloHatasi(ValueError):
    pass


def _num(v: Any) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return None


def _find_header_row(ws) -> Optional[int]:
    # Baslik satiri: A sutununda "BOY" gecen ve sagindaki hucreler EN
    # degerleri (artan sayilar) olan satir.
    for r in range(1, min(ws.max_row, 30) + 1):
        a = ws.cell(row=r, column=1).value
        if isinstance(a, str) and "BOY" in a.upper():
            nums = [_num(ws.cell(row=r, column=c).value) for c in range(2, ws.max_column + 1)]
            if sum(1 for n in nums if n) >= 3:
                return r
    return None


def _zam_pct(ws, header_row: int) -> float:
    # "Zam Oranı (%)" etiketinin sagindaki ilk sayisal hucre.
    for r in range(1, header_row):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and "ZAM" in v.upper():
                for c2 in range(c + 1, ws.max_column + 1):
                    n = _num(ws.cell(row=r, column=c2).value)
                    if n is not None:
                        return n
    return 0.0


def parse_excel(file_bytes: bytes) -> Dict[str, Any]:
    try:
        wb_val = load_workbook(io.BytesIO(file_bytes), data_only=True)
        wb_f = load_workbook(io.BytesIO(file_bytes), data_only=False)
    except Exception as e:
        raise ZipPerdeTabloHatasi(f"Excel dosyası okunamadı: {e}")

    for name in wb_f.sheetnames:
        ws_f, ws_v = wb_f[name], wb_val[name]
        hdr = _find_header_row(ws_f)
        if hdr is None:
            continue
        zam = _zam_pct(ws_v, hdr)
        cols: List[int] = []
        widths: List[int] = []
        for c in range(2, ws_f.max_column + 1):
            n = _num(ws_f.cell(row=hdr, column=c).value)
            if n and n > 0:
                cols.append(c)
                widths.append(int(round(n)))
        heights: List[int] = []
        prices: List[List[Optional[float]]] = []
        for r in range(hdr + 1, ws_f.max_row + 1):
            h = _num(ws_f.cell(row=r, column=1).value)
            if not h or h <= 0:
                continue
            row: List[Optional[float]] = []
            for c in cols:
                val = _num(ws_v.cell(row=r, column=c).value)
                if val is None:
                    raw = ws_f.cell(row=r, column=c).value
                    m = _ROUND_RE.match(raw) if isinstance(raw, str) else None
                    if m:
                        val = round(float(m.group(1)) * (1 + zam / 100.0))
                row.append(val if val and val > 0 else None)
            heights.append(int(round(h)))
            prices.append(row)
        if len(widths) < 3 or len(heights) < 3:
            continue
        if widths != sorted(widths) or heights != sorted(heights):
            raise ZipPerdeTabloHatasi("EN ve BOY değerleri küçükten büyüğe sıralı olmalı")
        filled = sum(1 for row in prices for p in row if p)
        if filled < 10:
            raise ZipPerdeTabloHatasi("Tabloda beklenenden az fiyat bulundu, dosyayı kontrol edin")
        return {"widths": widths, "heights": heights, "prices": prices, "currency": CURRENCY, "zamPct": zam}

    raise ZipPerdeTabloHatasi("Dosyada 'BOY↓ / EN→' başlıklı bir fiyat tablosu bulunamadı")


def lookup(table: Dict[str, Any], en_cm: float, boy_cm: float) -> Dict[str, Any]:
    widths, heights, prices = table["widths"], table["heights"], table["prices"]
    ci = next((i for i, w in enumerate(widths) if w >= en_cm), None)
    ri = next((i for i, h in enumerate(heights) if h >= boy_cm), None)
    if ci is None or ri is None or en_cm <= 0 or boy_cm <= 0:
        return {"ok": False, "reason": f"Ölçü tablo dışında (en fazla EN {widths[-1]} × BOY {heights[-1]} cm)"}
    price = prices[ri][ci]
    if not price:
        return {"ok": False, "reason": f"EN {widths[ci]} × BOY {heights[ri]} cm üretilmiyor"}
    return {"ok": True, "price": price, "en": widths[ci], "boy": heights[ri], "currency": table.get("currency", CURRENCY)}
