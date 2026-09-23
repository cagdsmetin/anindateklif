"""zip_perde birim testleri -- calisan backend gerektirmez.

Kapsam: paket icindeki varsayilan tablo, ust-basamak aramasi, uretilmeyen
hucreler, Excel ayristirma (hem onbellekli degerle hem formulden zam
uygulayarak) ve hatali dosyanin reddi.
"""
import io
import os
import sys

import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import zip_perde as zp  # noqa: E402

T = zp.DEFAULT_TABLE


def test_default_table_shape():
    assert T["currency"] == "EUR"
    assert T["widths"][0] == 125 and T["widths"][-1] == 580
    assert T["heights"][0] == 125 and T["heights"][-1] == 500
    assert all(len(row) == len(T["widths"]) for row in T["prices"])
    assert T["prices"][0][0] == 250


def test_lookup_exact_and_rounds_up():
    assert zp.lookup(T, 125, 125)["price"] == 250
    r = zp.lookup(T, 210, 240)
    assert (r["en"], r["boy"], r["price"]) == (225, 250, 330)


def test_lookup_out_of_range_and_not_produced():
    assert not zp.lookup(T, 600, 200)["ok"]
    assert not zp.lookup(T, 200, 520)["ok"]
    r = zp.lookup(T, 580, 500)
    assert not r["ok"] and "üretilmiyor" in r["reason"]


def _excel(zam=0, cached=True):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "SKYART"
    ws["A2"] = "Zam Oranı (%):"
    ws["D2"] = zam
    ws["A4"] = "BOY↓ / EN→"
    widths = [100, 200, 300, 400]
    for i, w in enumerate(widths):
        ws.cell(row=4, column=2 + i, value=w)
    for r, h in enumerate([100, 200, 300]):
        ws.cell(row=5 + r, column=1, value=h)
        for c in range(len(widths)):
            base = 100 + r * 10 + c
            val = base if cached else f"=ROUND({base}*(1+$D$2/100),0)"
            ws.cell(row=5 + r, column=2 + c, value=val)
    ws.cell(row=7, column=5, value="—")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_cached_values():
    t = zp.parse_excel(_excel())
    assert t["widths"] == [100, 200, 300, 400]
    assert t["heights"] == [100, 200, 300]
    assert t["prices"][0] == [100, 101, 102, 103]
    assert t["prices"][2][3] is None


def test_parse_formula_applies_zam():
    t = zp.parse_excel(_excel(zam=10, cached=False))
    assert t["zamPct"] == 10
    assert t["prices"][0][0] == 110
    assert t["prices"][1][2] == round(112 * 1.1)


def test_parse_rejects_unrelated_file():
    wb = Workbook()
    wb.active["A1"] = "merhaba"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(zp.ZipPerdeTabloHatasi):
        zp.parse_excel(buf.getvalue())
